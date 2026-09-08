"""Turn an uploaded file into ordered (page, paragraph) text blocks.

PDF and DOCX support are imported lazily so a stripped-down install still
starts and still handles plain text.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md"}

_WHITESPACE_RUN = re.compile(r"[ \t\r\f\v]+")
_BLANK_LINES = re.compile(r"\n\s*\n+")
_HYPHEN_LINEBREAK = re.compile(r"(\w)-\n(\w)")


class ExtractionError(RuntimeError):
    """Raised when a file cannot be turned into usable text."""


@dataclass(frozen=True)
class Block:
    """One paragraph of source text, with the page it came from (1-indexed)."""

    page: int | None
    text: str


def _clean(text: str) -> str:
    text = text.replace(" ", " ").replace("\x00", "")
    # Repair words hyphenated across a line break, a very common PDF artefact.
    text = _HYPHEN_LINEBREAK.sub(r"\1\2", text)
    text = _WHITESPACE_RUN.sub(" ", text)
    return text.strip()


def _paragraphs(page_text: str, page: int | None) -> list[Block]:
    cleaned = _clean(page_text)
    if not cleaned:
        return []
    blocks: list[Block] = []
    for para in _BLANK_LINES.split(cleaned):
        # Single newlines inside a paragraph are layout, not structure.
        merged = " ".join(line.strip() for line in para.split("\n") if line.strip())
        if merged:
            blocks.append(Block(page=page, text=merged))
    return blocks


def _read_text_file(path: Path) -> list[Block]:
    raw: str | None = None
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            raw = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raise ExtractionError("Could not decode the file as text.")

    # Form feeds are the only page signal a plain text file ever carries.
    paged = "\f" in raw
    blocks: list[Block] = []
    for page_number, page_text in enumerate(raw.split("\f"), start=1):
        blocks.extend(_paragraphs(page_text, page_number if paged else None))
    if not blocks:
        raise ExtractionError("The file is empty.")
    return blocks


def _read_pdf(path: Path) -> list[Block]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise ExtractionError(
            "PDF support requires the 'pypdf' package (pip install pypdf)."
        ) from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open the PDF: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")  # Many PDFs are "encrypted" with an empty password.
        except Exception as exc:
            raise ExtractionError("The PDF is password protected.") from exc

    blocks: list[Block] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        blocks.extend(_paragraphs(page_text, page_number))

    if not blocks:
        raise ExtractionError(
            "No selectable text found. The PDF is probably a scan; OCR is out of scope."
        )
    return blocks


def _read_docx(path: Path) -> list[Block]:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise ExtractionError(
            "DOCX support requires the 'python-docx' package (pip install python-docx)."
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open the DOCX: {exc}") from exc

    blocks: list[Block] = []
    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if text:
            blocks.append(Block(page=None, text=text))

    # Tables carry a lot of policy and specification content, so flatten them
    # row by row rather than dropping them.
    for table in document.tables:
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                blocks.append(Block(page=None, text=line))

    if not blocks:
        raise ExtractionError("The document contains no readable text.")
    return blocks


def extract(path: Path, *, extension: str | None = None) -> list[Block]:
    """Extract ordered text blocks from a PDF, TXT/MD or DOCX file."""
    suffix = (extension or path.suffix).lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix in (".txt", ".md"):
        return _read_text_file(path)
    if suffix == ".doc":
        raise ExtractionError("Legacy .doc files are not supported - save as .docx first.")
    raise ExtractionError(f"Unsupported file type '{suffix}'. Use PDF, TXT or DOCX.")
